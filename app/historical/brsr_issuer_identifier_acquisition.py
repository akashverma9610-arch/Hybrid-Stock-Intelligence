"""Exact issuer-identifier extraction from official NSE BRSR bulk artifacts.

The BRSR General workbook directly labels an NSE symbol, CIN, listed-entity
name, application identity, and submission timestamp.  This adapter preserves
that filing-level fact only.  It does not create a session bridge merely
because a later NSE session uses the same symbol: continuity is a separately
governed question.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from enum import StrEnum
import hashlib
from io import BytesIO
import json
import re
from typing import Iterable
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from app.historical.issuer_identifier_acquisition import (
    ExactIssuerIdentifierEvidence,
    IssuerIdentifierEvidenceSource,
    IssuerIdentifierEvidenceStore,
    IssuerIdentifierWriteOutcome,
    evidence_from_structured_row,
)
from app.historical.issuer_security_identity import IssuerIdentifierType


CONTRACT_ID = "brsr-primary-issuer-identifier-backfill"
CONTRACT_VERSION = "v1"
MAX_NSE_BRSR_ARCHIVE_BYTES = 64 * 1024 * 1024
MAX_NSE_BRSR_GENERAL_XLSX_BYTES = 16 * 1024 * 1024
_CIN = re.compile(r"[LU][0-9]{5}[A-Z]{2}[0-9]{4}[A-Z]{3}[0-9]{6}\Z")
_FP = re.compile(r"sha256:[0-9a-f]{64}\Z")
_FY = re.compile(r"FY[0-9]{2}-[0-9]{2}\Z")
_IST = ZoneInfo("Asia/Kolkata")


class BrsrIssuerIdentifierAcquisitionError(ValueError):
    pass


class BrsrTargetIssuerAuditStatus(StrEnum):
    EXACT_FILING_EVIDENCE_RETAINED = "EXACT_FILING_EVIDENCE_RETAINED"
    EVIDENCE_CONFLICTED = "EVIDENCE_CONFLICTED"
    NO_EXACT_BRSR_SYMBOL_EVIDENCE = "NO_EXACT_BRSR_SYMBOL_EVIDENCE"


@dataclass(frozen=True, slots=True)
class BrsrIssuerIdentifierExtraction:
    """One exact BRSR General-table extraction from a supplied artifact."""

    financial_year: str
    source_reference: str
    source_document_fingerprint: str
    workbook_member_path: str | None
    rows_inspected: int
    rows_rejected: int
    evidence: tuple[ExactIssuerIdentifierEvidence, ...]
    fingerprint: str
    contract_id: str = CONTRACT_ID
    contract_version: str = CONTRACT_VERSION

    def __post_init__(self) -> None:
        if not (
            _FY.fullmatch(self.financial_year)
            and _nonblank(self.source_reference)
            and _fingerprint(self.source_document_fingerprint)
            and (self.workbook_member_path is None or _safe_member_path(self.workbook_member_path))
            and self.rows_inspected >= 0
            and self.rows_rejected >= 0
            and self.rows_inspected == len(self.evidence) + self.rows_rejected
            and self.evidence == _canonical_evidence(self.evidence)
            and all(item.source is IssuerIdentifierEvidenceSource.NSE_BRSR_ARCHIVE for item in self.evidence)
            and all(item.isin is None and item.nse_symbol is not None for item in self.evidence)
            and self.contract_id == CONTRACT_ID
            and self.contract_version == CONTRACT_VERSION
            and self.fingerprint == _hash(_extraction_payload(self))
        ):
            raise BrsrIssuerIdentifierAcquisitionError("BRSR_ISSUER_IDENTIFIER_EXTRACTION_INVALID")


@dataclass(frozen=True, slots=True)
class BrsrIssuerIdentifierPersistenceReceipt:
    """Deterministic outcome of appending BRSR filing-level evidence only."""

    source_document_fingerprint: str
    evidence_inserted: int
    evidence_already_present: int
    fingerprint: str
    contract_id: str = CONTRACT_ID
    contract_version: str = CONTRACT_VERSION

    def __post_init__(self) -> None:
        if not (
            _fingerprint(self.source_document_fingerprint)
            and self.evidence_inserted >= 0
            and self.evidence_already_present >= 0
            and self.contract_id == CONTRACT_ID
            and self.contract_version == CONTRACT_VERSION
            and self.fingerprint == _hash(_persistence_payload(self))
        ):
            raise BrsrIssuerIdentifierAcquisitionError("BRSR_ISSUER_IDENTIFIER_PERSISTENCE_RECEIPT_INVALID")


@dataclass(frozen=True, slots=True)
class BrsrTargetIssuerAudit:
    """Read-only exact-symbol BRSR filing evidence for one named target."""

    symbol: str
    evidence: tuple[ExactIssuerIdentifierEvidence, ...]
    status: BrsrTargetIssuerAuditStatus

    def __post_init__(self) -> None:
        if not (
            _nonblank(self.symbol)
            and self.evidence == _canonical_evidence(self.evidence)
            and all(item.source is IssuerIdentifierEvidenceSource.NSE_BRSR_ARCHIVE and item.nse_symbol == self.symbol for item in self.evidence)
            and isinstance(self.status, BrsrTargetIssuerAuditStatus)
            and ((not self.evidence and self.status is BrsrTargetIssuerAuditStatus.NO_EXACT_BRSR_SYMBOL_EVIDENCE) or self.evidence)
            and (
                (not self.evidence and self.status is BrsrTargetIssuerAuditStatus.NO_EXACT_BRSR_SYMBOL_EVIDENCE)
                or (len({item.issuer_identifier_value for item in self.evidence}) == 1 and self.status is BrsrTargetIssuerAuditStatus.EXACT_FILING_EVIDENCE_RETAINED)
                or (len({item.issuer_identifier_value for item in self.evidence}) > 1 and self.status is BrsrTargetIssuerAuditStatus.EVIDENCE_CONFLICTED)
            )
        ):
            raise BrsrIssuerIdentifierAcquisitionError("BRSR_TARGET_ISSUER_AUDIT_INVALID")


def extract_nse_brsr_general_xlsx(
    *, payload: bytes, source_reference: str, financial_year: str, source_observed_at: datetime,
    workbook_member_path: str | None = None,
) -> BrsrIssuerIdentifierExtraction:
    """Extract only direct labelled BRSR General-table identity fields.

    The parser is intentionally not a document-understanding layer: one
    unambiguous General sheet must provide all linkage fields in the same row.
    """
    _validate_source_metadata(payload, source_reference, financial_year, source_observed_at, workbook_member_path)
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # Consistent with the existing NSE import path.
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_XLSX_PARSER_UNAVAILABLE") from error
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_vba=False)
        candidates = tuple(_sheet_header(sheet) for sheet in workbook.worksheets)
        matches = tuple(item for item in candidates if _is_general_identity_header(item[1]))
        if len(matches) != 1:
            raise BrsrIssuerIdentifierAcquisitionError("BRSR_GENERAL_IDENTITY_SHEET_AMBIGUOUS")
        sheet, header = matches[0]
        rows = _extract_rows(
            sheet=sheet, header=header, source_document_fingerprint=_bytes_hash(payload),
            source_reference=_member_reference(source_reference, workbook_member_path),
            source_observed_at=_utc(source_observed_at),
        )
    except BrsrIssuerIdentifierAcquisitionError:
        raise
    except (KeyError, OSError, ValueError, TypeError) as error:
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_GENERAL_XLSX_INVALID") from error
    evidence, rejected = rows
    values = dict(
        financial_year=financial_year, source_reference=source_reference,
        source_document_fingerprint=_bytes_hash(payload), workbook_member_path=workbook_member_path,
        rows_inspected=len(evidence) + rejected, rows_rejected=rejected, evidence=evidence,
    )
    return BrsrIssuerIdentifierExtraction(
        **values, fingerprint=_hash(_extraction_payload_from_values(values)),
    )


def extract_nse_brsr_general_zip(
    *, payload: bytes, source_reference: str, financial_year: str, source_observed_at: datetime,
    workbook_member_path: str,
) -> BrsrIssuerIdentifierExtraction:
    """Extract one explicitly named BRSR General workbook from a bulk ZIP.

    The caller must supply the exact member path discovered from NSE; this
    function never pattern-selects a workbook from a changing archive.
    """
    if not payload or len(payload) > MAX_NSE_BRSR_ARCHIVE_BYTES:
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_ARCHIVE_SIZE_INVALID")
    if not _safe_member_path(workbook_member_path):
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_ARCHIVE_MEMBER_PATH_INVALID")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            info = archive.getinfo(workbook_member_path)
            if info.is_dir() or info.file_size > MAX_NSE_BRSR_GENERAL_XLSX_BYTES:
                raise BrsrIssuerIdentifierAcquisitionError("BRSR_ARCHIVE_MEMBER_SIZE_INVALID")
            workbook = archive.read(info)
    except BrsrIssuerIdentifierAcquisitionError:
        raise
    except (BadZipFile, KeyError, OSError) as error:
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_ARCHIVE_INVALID") from error
    return extract_nse_brsr_general_xlsx(
        payload=workbook, source_reference=source_reference, financial_year=financial_year,
        source_observed_at=source_observed_at, workbook_member_path=workbook_member_path,
    )


def persist_brsr_issuer_identifier_evidence(
    *, store: IssuerIdentifierEvidenceStore, extraction: BrsrIssuerIdentifierExtraction,
) -> BrsrIssuerIdentifierPersistenceReceipt:
    """Append filing-level BRSR evidence, deliberately creating no bridge."""
    if not isinstance(store, IssuerIdentifierEvidenceStore) or not isinstance(extraction, BrsrIssuerIdentifierExtraction):
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_EVIDENCE_PERSISTENCE_INPUT_INVALID")
    inserted = present = 0
    for item in extraction.evidence:
        outcome = store.save_evidence(item)
        inserted += outcome is IssuerIdentifierWriteOutcome.INSERTED
        present += outcome is IssuerIdentifierWriteOutcome.ALREADY_PRESENT
    values = dict(
        source_document_fingerprint=extraction.source_document_fingerprint,
        evidence_inserted=inserted, evidence_already_present=present,
    )
    return BrsrIssuerIdentifierPersistenceReceipt(
        **values, fingerprint=_hash(_persistence_payload_from_values(values)),
    )


def audit_brsr_target_evidence(
    *, evidence: Iterable[ExactIssuerIdentifierEvidence], targets: Iterable[str],
) -> tuple[BrsrTargetIssuerAudit, ...]:
    """Audit exact BRSR symbols only; names never participate in the lookup."""
    values = _canonical_evidence(evidence)
    audits: list[BrsrTargetIssuerAudit] = []
    for symbol in tuple(sorted(set(targets))):
        if not _nonblank(symbol):
            raise BrsrIssuerIdentifierAcquisitionError("BRSR_TARGET_SYMBOL_INVALID")
        matches = tuple(item for item in values if item.source is IssuerIdentifierEvidenceSource.NSE_BRSR_ARCHIVE and item.nse_symbol == symbol)
        cins = {item.issuer_identifier_value for item in matches}
        status = BrsrTargetIssuerAuditStatus.NO_EXACT_BRSR_SYMBOL_EVIDENCE if not matches else (
            BrsrTargetIssuerAuditStatus.EVIDENCE_CONFLICTED if len(cins) > 1 else BrsrTargetIssuerAuditStatus.EXACT_FILING_EVIDENCE_RETAINED
        )
        audits.append(BrsrTargetIssuerAudit(symbol=symbol, evidence=matches, status=status))
    return tuple(audits)


def _extract_rows(*, sheet: object, header: dict[str, int], source_document_fingerprint: str,
                  source_reference: str, source_observed_at: datetime) -> tuple[tuple[ExactIssuerIdentifierEvidence, ...], int]:
    aliases = {
        "application": ("APP_ID", "APPLICATION_NUMBER", "BRSR_SUBMISSION_DTLS_ID"),
        "symbol": ("SYMB_SYMBOL", "NSE SYMBOL", "SYMBOL"),
        "cin": ("CORPORATE IDENTITY NUMBER (CIN) OF THE LISTED ENTITY", "CIN", "COMPANY_CIN"),
        "submitted": ("TLA_SUBMITTED_DT", "DATE OF SUBMISSION", "PDF_SUBMISSION_DT"),
        "legal_name": ("NAME OF THE LISTED ENTITY", "COMPANY_NAME", "NAME OF THE COMPANY"),
    }
    indexes = {name: _required_column(header, names) for name, names in aliases.items() if name != "legal_name"}
    # BRSR FY23-24 exposes both a source-system company display name and the
    # specifically labelled listed-entity name.  The latter is the exact
    # listed-entity field; the former is a fallback only when it is absent.
    legal_index = _preferred_column(
        header, "NAME OF THE LISTED ENTITY", ("COMPANY_NAME", "NAME OF THE COMPANY"),
    )
    values: list[ExactIssuerIdentifierEvidence] = []
    rejected = 0
    by_record_id: dict[str, ExactIssuerIdentifierEvidence] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):  # type: ignore[union-attr]
        application = _application_value(row, indexes["application"])
        symbol = _value(row, indexes["symbol"])
        cin = _value(row, indexes["cin"])
        submitted = _value(row, indexes["submitted"])
        if not any((application, symbol, cin, submitted)):
            continue
        if not (_nonblank(application) and _nonblank(symbol) and _CIN.fullmatch(cin or "")):
            rejected += 1
            continue
        known_at = _submission_time(submitted)
        record_id = "nse-brsr-application:" + application
        item = evidence_from_structured_row(
            source=IssuerIdentifierEvidenceSource.NSE_BRSR_ARCHIVE,
            source_document_fingerprint=source_document_fingerprint, source_reference=source_reference,
            source_record_id=record_id, source_as_of_date=known_at.astimezone(_IST).date(),
            source_observed_at=source_observed_at, historical_known_at=known_at, isin=None,
            nse_symbol=symbol, issuer_identifier_type=IssuerIdentifierType.CIN,
            issuer_identifier_value=cin, issuer_legal_name=_value(row, legal_index) if legal_index is not None else None,
        )
        prior = by_record_id.get(record_id)
        if prior is not None:
            if prior != item:
                raise BrsrIssuerIdentifierAcquisitionError("BRSR_FILING_ROW_AMBIGUOUS")
            continue
        by_record_id[record_id] = item
        values.append(item)
    return _canonical_evidence(values), rejected


def _sheet_header(sheet: object) -> tuple[object, dict[str, int]]:
    row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())  # type: ignore[union-attr]
    header: dict[str, int] = {}
    for index, value in enumerate(row):
        name = _header(value)
        if not name:
            continue
        if name in header:
            raise BrsrIssuerIdentifierAcquisitionError("BRSR_GENERAL_HEADER_AMBIGUOUS")
        header[name] = index
    return sheet, header


def _is_general_identity_header(header: dict[str, int]) -> bool:
    return all(_optional_column(header, choices) is not None for choices in (
        ("APP_ID", "APPLICATION_NUMBER", "BRSR_SUBMISSION_DTLS_ID"),
        ("SYMB_SYMBOL", "NSE SYMBOL", "SYMBOL"),
        ("CORPORATE IDENTITY NUMBER (CIN) OF THE LISTED ENTITY", "CIN", "COMPANY_CIN"),
        ("TLA_SUBMITTED_DT", "DATE OF SUBMISSION", "PDF_SUBMISSION_DT"),
    ))


def _required_column(header: dict[str, int], names: tuple[str, ...]) -> int:
    value = _optional_column(header, names)
    if value is None:
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_GENERAL_IDENTITY_HEADERS_MISSING")
    return value


def _optional_column(header: dict[str, int], names: tuple[str, ...]) -> int | None:
    matches = tuple(header[name] for name in names if name in header)
    if len(matches) > 1:
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_GENERAL_IDENTITY_HEADER_ALIAS_AMBIGUOUS")
    return matches[0] if matches else None


def _preferred_column(header: dict[str, int], preferred: str, fallbacks: tuple[str, ...]) -> int | None:
    return header[preferred] if preferred in header else _optional_column(header, fallbacks)


def _validate_source_metadata(payload: bytes, source_reference: str, financial_year: str,
                              source_observed_at: datetime, member_path: str | None) -> None:
    if not payload or len(payload) > MAX_NSE_BRSR_GENERAL_XLSX_BYTES:
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_GENERAL_XLSX_SIZE_INVALID")
    if not _nonblank(source_reference) or _FY.fullmatch(financial_year) is None or not _aware(source_observed_at):
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_SOURCE_METADATA_INVALID")
    if member_path is not None and not _safe_member_path(member_path):
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_ARCHIVE_MEMBER_PATH_INVALID")


def _submission_time(value: object) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        text = value.strip()
        parsed = None
        for format_ in ("%d-%b-%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
            try:
                parsed = datetime.strptime(text.upper() if "%b" in format_ else text, format_)
                break
            except ValueError:
                continue
        if parsed is None:
            raise BrsrIssuerIdentifierAcquisitionError("BRSR_FILING_TIMESTAMP_INVALID")
    else:
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_FILING_TIMESTAMP_INVALID")
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        parsed = parsed.replace(tzinfo=_IST)
    return parsed.astimezone(UTC)


def _canonical_evidence(values: Iterable[ExactIssuerIdentifierEvidence]) -> tuple[ExactIssuerIdentifierEvidence, ...]:
    supplied = tuple(values)
    if not all(isinstance(item, ExactIssuerIdentifierEvidence) for item in supplied):
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_EVIDENCE_INPUT_INVALID")
    unique: dict[tuple[str, str], ExactIssuerIdentifierEvidence] = {}
    for item in supplied:
        item.verify()
        key = (item.evidence_id, item.fingerprint)
        prior = unique.get(key)
        if prior is not None and prior != item:
            raise BrsrIssuerIdentifierAcquisitionError("BRSR_EVIDENCE_DUPLICATE_CONFLICT")
        unique[key] = item
    return tuple(sorted(unique.values(), key=lambda item: (item.source_as_of_date, item.evidence_id)))


def _extraction_payload(value: BrsrIssuerIdentifierExtraction) -> dict[str, object]:
    return _extraction_payload_from_values({name: getattr(value, name) for name in value.__dataclass_fields__ if name not in {"fingerprint", "contract_id", "contract_version"}})


def _extraction_payload_from_values(value: dict[str, object]) -> dict[str, object]:
    return {
        "contract_id": CONTRACT_ID, "contract_version": CONTRACT_VERSION,
        "financial_year": value["financial_year"], "source_reference": value["source_reference"],
        "source_document_fingerprint": value["source_document_fingerprint"], "workbook_member_path": value["workbook_member_path"],
        "rows_inspected": value["rows_inspected"], "rows_rejected": value["rows_rejected"],
        "evidence": tuple(item.fingerprint for item in value["evidence"]),
    }


def _persistence_payload(value: BrsrIssuerIdentifierPersistenceReceipt) -> dict[str, object]:
    return _persistence_payload_from_values({name: getattr(value, name) for name in value.__dataclass_fields__ if name not in {"fingerprint", "contract_id", "contract_version"}})


def _persistence_payload_from_values(value: dict[str, object]) -> dict[str, object]:
    return {
        "contract_id": CONTRACT_ID, "contract_version": CONTRACT_VERSION,
        "source_document_fingerprint": value["source_document_fingerprint"],
        "evidence_inserted": value["evidence_inserted"], "evidence_already_present": value["evidence_already_present"],
    }


def _value(row: tuple[object, ...], index: int | None) -> str | None:
    if index is None or index >= len(row):
        return None
    value = row[index]
    return value.strip() if isinstance(value, str) and value.strip() else None


def _application_value(row: tuple[object, ...], index: int) -> str | None:
    if index >= len(row):
        return None
    value = row[index]
    if isinstance(value, str):
        return value.strip() or None
    # NSE's BRSR General export serializes APP_ID as an XLSX numeric cell.  An
    # integer is an exact application identifier; floats, booleans, and other
    # display conversions are deliberately not accepted.
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return None


def _header(value: object) -> str:
    return " ".join(str(value or "").strip().upper().split())


def _member_reference(source_reference: str, path: str | None) -> str:
    return source_reference if path is None else source_reference + "#" + path


def _safe_member_path(value: object) -> bool:
    return isinstance(value, str) and value.endswith(".xlsx") and not value.startswith(("/", "\\")) and ".." not in value.split("/") and "\\" not in value


def _nonblank(value: object) -> bool:
    return isinstance(value, str) and bool(value.strip())


def _aware(value: object) -> bool:
    return isinstance(value, datetime) and value.tzinfo is not None and value.utcoffset() is not None


def _utc(value: datetime) -> datetime:
    if not _aware(value):
        raise BrsrIssuerIdentifierAcquisitionError("BRSR_AWARE_TIMESTAMP_REQUIRED")
    return value.astimezone(UTC)


def _fingerprint(value: object) -> bool:
    return isinstance(value, str) and _FP.fullmatch(value) is not None


def _bytes_hash(value: bytes) -> str:
    return "sha256:" + hashlib.sha256(value).hexdigest()


def _hash(value: object) -> str:
    return "sha256:" + hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode(),
    ).hexdigest()


__all__ = (
    "CONTRACT_ID", "CONTRACT_VERSION", "MAX_NSE_BRSR_ARCHIVE_BYTES", "MAX_NSE_BRSR_GENERAL_XLSX_BYTES",
    "BrsrIssuerIdentifierAcquisitionError", "BrsrIssuerIdentifierExtraction",
    "BrsrIssuerIdentifierPersistenceReceipt", "BrsrTargetIssuerAudit", "BrsrTargetIssuerAuditStatus", "audit_brsr_target_evidence",
    "extract_nse_brsr_general_xlsx", "extract_nse_brsr_general_zip",
    "persist_brsr_issuer_identifier_evidence",
)
